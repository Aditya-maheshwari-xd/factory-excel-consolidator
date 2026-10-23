"""Build the one summary the owner actually reads, as Excel and as HTML."""

from __future__ import annotations

import html
import os
from collections import Counter
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .checks import ERROR, INFO, WARNING, Issue

def plant_now() -> datetime:
    """Current time at the plant, not on the server (servers often run in UTC)."""
    return datetime.now(ZoneInfo(os.environ.get("PLANT_TZ", "Asia/Kolkata")))


CLEAN_COLUMNS = ["date", "shift", "machine", "part", "operator", "target_qty", "produced_qty",
                 "rejected_qty", "downtime_min", "remarks", "counted", "source_file", "source_sheet",
                 "source_row"]


def to_frame(rows: list[dict]) -> pd.DataFrame:
    df = pd.DataFrame(rows, columns=CLEAN_COLUMNS)
    for c in ("target_qty", "produced_qty", "rejected_qty", "downtime_min"):
        df[c] = pd.to_numeric(df[c], errors="coerce")
    return df


def _rollup(df: pd.DataFrame, by: list[str]) -> pd.DataFrame:
    placed = df[df["counted"].fillna(True).astype(bool)].dropna(subset=by)
    if placed.empty:
        return pd.DataFrame(columns=by + ["target", "produced", "rejected", "ok_qty",
                                          "achievement_pct", "rejection_pct", "downtime_min"])
    g = placed.groupby(by, dropna=False).agg(
        target=("target_qty", "sum"), produced=("produced_qty", "sum"),
        rejected=("rejected_qty", "sum"), downtime_min=("downtime_min", "sum")).reset_index()
    g["ok_qty"] = g["produced"] - g["rejected"]
    g["achievement_pct"] = (g["produced"] / g["target"].where(g["target"] > 0)).round(3)
    g["rejection_pct"] = (g["rejected"] / g["produced"].where(g["produced"] > 0)).round(3)
    return g[by + ["target", "produced", "rejected", "ok_qty", "achievement_pct",
                   "rejection_pct", "downtime_min"]]


def build_summaries(rows: list[dict]) -> dict[str, pd.DataFrame]:
    df = to_frame(rows)
    return {
        "daily": _rollup(df, ["date"]),
        "by_shift": _rollup(df, ["date", "shift"]),
        "by_machine": _rollup(df, ["machine"]),
        "clean": df,
    }


ISSUE_COLUMNS = ["severity", "area", "date", "shift", "machine", "item", "problem", "where_to_fix", "check"]


def issues_frame(issues: list[Issue]) -> pd.DataFrame:
    return pd.DataFrame([{
        "severity": i.severity.upper(), "area": i.area, "date": i.date, "shift": i.shift,
        "machine": i.machine, "item": i.item, "problem": i.message, "where_to_fix": i.where(),
        "check": i.check,
    } for i in issues], columns=ISSUE_COLUMNS)


# --------------------------------------------------------------------------
# Excel
# --------------------------------------------------------------------------

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
SEV_FILL = {"ERROR": PatternFill("solid", fgColor="F8D7DA"),
            "WARNING": PatternFill("solid", fgColor="FFF3CD"),
            "INFO": PatternFill("solid", fgColor="E2E3E5")}
PCT_COLS = {"achievement_pct", "rejection_pct"}
DATE_COLS = {"date", "due_date"}


def _write_sheet(wb: Workbook, title: str, df: pd.DataFrame) -> None:
    ws = wb.create_sheet(title)
    headers = [c.replace("_", " ").title() for c in df.columns]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center")
    for record in df.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and pd.isna(v)) else v for v in record])
    for idx, col in enumerate(df.columns, start=1):
        letter = get_column_letter(idx)
        width = max([len(headers[idx - 1])] + [len(str(c.value or "")) for c in ws[letter][1:200]])
        ws.column_dimensions[letter].width = min(max(width + 2, 9), 70)
        for c in ws[letter][1:]:
            if col in PCT_COLS:
                c.number_format = "0.0%"
            elif col in DATE_COLS:
                c.number_format = "DD-MMM-YYYY"
    if "severity" in df.columns:
        for row in ws.iter_rows(min_row=2):
            fill = SEV_FILL.get(row[0].value)
            if fill:
                for c in row:
                    c.fill = fill
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def write_excel(path: str | Path, summaries: dict[str, pd.DataFrame], issues: list[Issue],
                stock: pd.DataFrame | None = None, dispatch: pd.DataFrame | None = None) -> Path:
    path = Path(path)
    wb = Workbook()
    wb.remove(wb.active)
    _write_sheet(wb, "Daily Summary", summaries["daily"])
    _write_sheet(wb, "By Shift", summaries["by_shift"])
    _write_sheet(wb, "By Machine", summaries["by_machine"])
    _write_sheet(wb, "Issues To Fix", issues_frame(issues))
    _write_sheet(wb, "All Entries (Clean)", summaries["clean"])
    if stock is not None and not stock.empty:
        _write_sheet(wb, "Stock Summary", stock)
    if dispatch is not None and not dispatch.empty:
        _write_sheet(wb, "Dispatch Summary", dispatch)
    wb.save(path)
    return path


# --------------------------------------------------------------------------
# HTML
# --------------------------------------------------------------------------

def _fmt(v: object, col: str) -> str:
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return "—"
    if col in PCT_COLS:
        return f"{float(v):.1%}"
    if isinstance(v, float):
        # Pieces are whole numbers; kilograms often aren't (5.88 kg must not show as 6).
        return f"{v:,.0f}" if v.is_integer() else f"{v:,.2f}"
    if hasattr(v, "strftime"):
        return v.strftime("%d %b %Y")
    return html.escape(str(v))


def _table(df: pd.DataFrame, row_class: str | None = None, drop_empty: bool = False) -> str:
    if drop_empty and not df.empty:      # e.g. no Item column when only production sheets were read
        df = df[[c for c in df.columns if df[c].notna().any() or c == row_class]]
    head = "".join(f"<th>{html.escape(c.replace('_', ' ').title())}</th>" for c in df.columns)
    body = []
    for rec in df.to_dict("records"):
        cls = f' class="{str(rec.get(row_class, "")).lower()}"' if row_class else ""
        body.append(f"<tr{cls}>" + "".join(f'<td class="c-{c}">{_fmt(rec[c], c)}</td>' for c in df.columns)
                    + "</tr>")
    return f"<table><thead><tr>{head}</tr></thead><tbody>{''.join(body)}</tbody></table>"


def write_html(path: str | Path, summaries: dict[str, pd.DataFrame], issues: list[Issue],
               plant_name: str = "Plant", files_read: int = 0, header_html: str = "",
               stock: pd.DataFrame | None = None, dispatch: pd.DataFrame | None = None,
               as_of=None) -> Path:
    """header_html: optional block shown above the report (the web page adds its download bar here).
    stock / dispatch: summaries of stock and dispatch registers, shown when those sheets were read."""
    path = Path(path)
    stock = stock if stock is not None else pd.DataFrame()
    dispatch = dispatch if dispatch is not None else pd.DataFrame()
    daily = summaries["daily"]
    counts = Counter(i.severity for i in issues)
    tiles = []
    if not daily.empty:
        total_prod, total_rej, total_tgt = daily["produced"].sum(), daily["rejected"].sum(), daily["target"].sum()
        tiles += [
            ("Produced", f"{total_prod:,.0f}"),
            ("Achievement", f"{total_prod / total_tgt:.1%}" if total_tgt else "—"),
            ("Rejection", f"{total_rej / total_prod:.1%}" if total_prod else "—"),
        ]
    if not stock.empty:
        tiles.append(("Stock items low or negative", str(int(stock["status"].isin(["low", "negative"]).sum()))))
    if not dispatch.empty:
        tiles += [("Orders overdue", str(int((dispatch["days_overdue"] > 0).sum()))),
                  ("Trucks awaiting POD", str(int(dispatch["awaiting_pod"].sum())))]
    tiles += [("Errors to fix", str(counts[ERROR])), ("Warnings", str(counts[WARNING]))]
    tiles_html = "".join(f'<div class="tile"><span>{k}</span><b>{v}</b></div>' for k, v in tiles)

    shown = issues_frame([i for i in issues if i.severity != INFO])
    clean = summaries["clean"]
    excluded = int((~clean["counted"].fillna(True).astype(bool)).sum()) if not clean.empty else 0
    excluded_note = (f"{excluded} entr{'y' if excluded == 1 else 'ies'} with impossible figures "
                     f"left out of totals — see Issues.") if excluded else ""
    only_production = stock.empty and dispatch.empty
    title = "Production Summary" if only_production else "Daily Factory Summary"
    as_of_note = (f" Overdue and POD checks are as of {as_of:%d %b %Y}, the latest date in the sheets."
                  if as_of and not dispatch.empty else "")

    sections = []
    if not daily.empty:
        sections += [f"<h2>Production: daily</h2><div class=\"wrap\">{_table(daily)}</div>",
                     f"<h2>Production: by shift</h2><div class=\"wrap\">{_table(summaries['by_shift'])}</div>",
                     f"<h2>Production: by machine</h2><div class=\"wrap\">{_table(summaries['by_machine'])}</div>"]
    if not stock.empty:
        sections.append(f"<h2>Stock register</h2><div class=\"wrap\">{_table(stock, row_class='status')}</div>")
    if not dispatch.empty:
        sections.append(f"<h2>Dispatch register: by order</h2><div class=\"wrap\">"
                        f"{_table(dispatch, row_class='status')}</div>")

    doc = f"""<!doctype html><html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>{html.escape(plant_name)} {title}</title>
<style>
:root{{--bg:#f6f7f9;--card:#fff;--ink:#1c2430;--muted:#5b6675;--line:#e3e6ea;--accent:#1f3a5f;
--err:#f8d7da;--warn:#fff3cd}}
@media (prefers-color-scheme:dark){{:root{{--bg:#12161c;--card:#1b212a;--ink:#e7ebf0;--muted:#9aa5b4;
--line:#2c3440;--accent:#8fb4e3;--err:#5a2a2f;--warn:#4d4220}}}}
body{{margin:0;background:var(--bg);color:var(--ink);font:14px/1.5 system-ui,sans-serif}}
main{{max-width:1100px;margin:0 auto;padding:24px 16px}}
h1{{margin:0 0 4px;font-size:22px}} h2{{margin:28px 0 8px;font-size:16px;color:var(--accent)}}
a{{color:var(--accent)}} .muted{{color:var(--muted)}}
.tiles{{display:grid;grid-template-columns:repeat(auto-fit,minmax(150px,1fr));gap:10px;margin-top:16px}}
.tile{{background:var(--card);border:1px solid var(--line);border-radius:8px;padding:12px}}
.tile span{{display:block;color:var(--muted);font-size:12px}} .tile b{{font-size:22px}}
.wrap{{overflow-x:auto;background:var(--card);border:1px solid var(--line);border-radius:8px}}
table{{border-collapse:collapse;width:100%}} th,td{{padding:7px 10px;border-bottom:1px solid var(--line);
text-align:left;white-space:nowrap}} th{{font-size:12px;color:var(--muted);font-weight:600}}
tr.error td,tr.negative td,tr.over-sent td{{background:var(--err)}}
tr.warning td,tr.low td{{background:var(--warn)}}
td.c-problem{{white-space:normal;min-width:280px}} td.c-where_to_fix{{white-space:normal;min-width:180px}}
</style></head><body><main>
{header_html}
<h1>{html.escape(plant_name)} — {title}</h1>
<div class="muted">Built from {files_read} file(s) on {plant_now():%d %b %Y, %H:%M}. {excluded_note}{as_of_note}</div>
<div class="tiles">{tiles_html}</div>
<h2>Issues to fix ({len(shown)})</h2><div class="wrap">{_table(shown, row_class="severity", drop_empty=True)}</div>
{"".join(sections)}
</main></body></html>"""
    path.write_text(doc, encoding="utf-8")
    return path
