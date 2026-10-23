"""The upload page: same results as the command line, and safe with bad uploads."""

import io
import re
from datetime import timedelta

import openpyxl
import pytest
from fastapi.testclient import TestClient

from consolidator import web
from consolidator.report import plant_now


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(web, "JOBS_DIR", tmp_path / "jobs")
    return TestClient(web.app)


def sample_uploads():
    return [("files", (name, (web.SAMPLE_DIR / name).read_bytes(),
                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"))
            for name in web._sample_files()]


def xlsx_bytes(rows):
    wb = openpyxl.Workbook()
    for r in rows:
        wb.active.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_home_lists_sample_sheets(client):
    r = client.get("/")
    assert r.status_code == 200
    for name in web._sample_files():
        assert name in r.text


def test_sample_run_matches_command_line(client):
    r = client.get("/sample")
    assert r.status_code == 200
    # 8 errors + 12 warnings, the same as `python -m consolidator` on the sample data
    assert "Issues to fix (20)" in r.text
    assert "Download the Excel report" in r.text


def test_upload_gives_same_result_and_excel_download(client):
    r = client.post("/consolidate", files=sample_uploads(), data={"use_sample_config": "1"})
    assert r.status_code == 200 and "Issues to fix (20)" in r.text
    job = re.search(r"/download/([0-9a-f]{32})", r.text).group(1)
    xl = client.get(f"/download/{job}")
    assert xl.status_code == 200 and xl.content[:2] == b"PK"      # an .xlsx is a zip file
    wb = openpyxl.load_workbook(io.BytesIO(xl.content))
    assert {"Issues To Fix", "Stock Summary", "Dispatch Summary"} <= set(wb.sheetnames)


def test_without_machine_list_capacity_and_missing_checks_are_off(client):
    r = client.post("/consolidate", files=sample_uploads())
    assert r.status_code == 200
    assert "No entry for" not in r.text          # missing-entry check needs the machine list


def test_non_excel_file_is_refused(client):
    r = client.post("/consolidate", files=[("files", ("notes.txt", b"hello", "text/plain"))])
    assert r.status_code == 400 and "Only Excel files" in r.text


def test_sheet_without_production_table_explains_why(client):
    data = xlsx_bytes([["Name", "City"], ["Ravi", "Pune"]])
    r = client.post("/consolidate", files=[("files", ("contacts.xlsx", data, "application/octet-stream"))])
    assert r.status_code == 422 and "No factory data found" in r.text
    assert "contacts.xlsx" in r.text


def test_corrupt_excel_does_not_crash(client):
    r = client.post("/consolidate", files=[("files", ("broken.xlsx", b"not really excel", "application/octet-stream"))])
    assert r.status_code == 422 and "could not be read" in r.text


def test_too_large_file_is_refused(client, monkeypatch):
    monkeypatch.setattr(web, "MAX_FILE_BYTES", 10)
    r = client.post("/consolidate", files=[("files", ("big.xlsx", b"x" * 11, "application/octet-stream"))])
    assert r.status_code == 413


def test_only_real_reports_and_sample_files_can_be_downloaded(client):
    assert client.get("/download/not-a-job").status_code == 404
    assert client.get("/download/" + "0" * 32).status_code == 404
    assert client.get("/samples/plant_config.example.json").status_code == 404
    assert client.get("/samples/" + web._sample_files()[0]).status_code == 200


def test_file_names_keep_what_the_checks_need():
    assert web._safe_name("press_shop_15-10-2026_revised.xlsx") == "press_shop_15-10-2026_revised.xlsx"
    assert web._safe_name("report (1).xlsx") == "report (1).xlsx"
    assert web._safe_name("../../etc/passwd.xlsx") == "passwd.xlsx"
    assert web._safe_name("C:\\Users\\x\\day<1>.xlsx").endswith("day_1_.xlsx")


def test_report_time_is_plant_time(monkeypatch):
    monkeypatch.setenv("PLANT_TZ", "Asia/Kolkata")
    assert plant_now().utcoffset() == timedelta(hours=5, minutes=30)
