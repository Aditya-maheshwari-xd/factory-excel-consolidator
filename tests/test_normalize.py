from datetime import date

import pytest
from openpyxl import Workbook

from consolidator.normalize import (match_header, normalize_machine, normalize_shift,
                                    parse_date, parse_qty, read_workbook)


@pytest.mark.parametrize("header,expected", [
    ("Prod. Qty", "produced_qty"),
    ("OUTPUT (pcs)", "produced_qty"),
    ("Rej. Qty", "rejected_qty"),
    ("M/C No", "machine"),
    ("Breakdown (min)", "downtime_min"),
    ("Downtime in mins", "downtime_min"),
    ("Item Code", "part"),
    ("Sr No", None),
    (None, None),
])
def test_match_header(header, expected):
    assert match_header(header) == expected


@pytest.mark.parametrize("raw,expected", [
    ("1,200", (1200.0, True)),
    ("395 pcs", (395.0, True)),
    ("1200 nos", (1200.0, True)),
    ("nil", (0.0, True)),
    ("-", (0.0, True)),
    (None, (None, True)),
    ("NA", (None, True)),
    (42, (42.0, True)),
    ("approx 750", (None, False)),
])
def test_parse_qty(raw, expected):
    assert parse_qty(raw) == expected


def test_parse_date_is_day_first():
    # 03/04/2026 on an Indian shop floor is 3rd April, never 4th March.
    assert parse_date("03/04/2026") == date(2026, 4, 3)
    assert parse_date("Date: 14-10-26") == date(2026, 10, 14)
    assert parse_date("cnc_2026-10-14") == date(2026, 10, 14)
    assert parse_date(46309) == date(2026, 10, 14)  # Excel serial
    assert parse_date("31/02/2026") is None          # impossible date is rejected, not guessed
    assert parse_date("Sheet1") is None


@pytest.mark.parametrize("raw,expected", [
    ("1st", "A"), ("Day", "A"), ("II", "B"), ("Night", "C"), (2, "B"), ("Shift B", "B"), ("x", None),
])
def test_normalize_shift(raw, expected):
    assert normalize_shift(raw) == expected


@pytest.mark.parametrize("raw,expected", [
    ("M-1", "M01"), ("m01", "M01"), ("Machine 12", "M12"), (3, "M03"), ("MC 4", "M04"),
])
def test_normalize_machine(raw, expected):
    assert normalize_machine(raw) == expected


def test_reads_sheet_with_banner_and_total_row(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Production"
    ws.append(["ACME FORGINGS DAILY REPORT"])
    ws.append(["Date: 14/10/2026"])
    ws.append([])
    ws.append(["M/C No", "Shift", "Target", "Prod. Qty", "Rej. Qty"])
    ws.append(["M-1", "1st", 1000, "980", 5])
    ws.append(["M-2", "2nd", 1000, "", 0])
    ws.append([])
    ws.append(["TOTAL", None, 2000, 980, 5])
    path = tmp_path / "report.xlsx"
    wb.save(path)

    result = read_workbook(path)
    assert len(result.rows) == 2  # total row and blank line ignored
    first = result.rows[0]
    assert first["date"] == date(2026, 10, 14)  # taken from the banner
    assert (first["machine"], first["shift"], first["produced_qty"]) == ("M01", "A", 980.0)
    assert first["source_row"] == 5  # row number as the supervisor sees it in Excel
    assert result.rows[1]["produced_qty"] is None


def test_shift_and_date_from_sheet_and_file_name(tmp_path):
    wb = Workbook()
    wb.active.title = "Notes"
    wb.active.append(["fill by 9am"])
    ws = wb.create_sheet("Shift B")
    ws.append(["Machine", "Plan", "Output", "Rejection"])
    ws.append(["M04", 400, "380 pcs", "nil"])
    path = tmp_path / "cnc_2026-10-15.xlsx"
    wb.save(path)

    result = read_workbook(path)
    assert result.sheets_skipped == ["cnc_2026-10-15.xlsx / Notes (no production table found)"]
    row = result.rows[0]
    assert (row["date"], row["shift"], row["produced_qty"], row["rejected_qty"]) == \
        (date(2026, 10, 15), "B", 380.0, 0.0)


def test_unopenable_file_is_reported_not_crashed(tmp_path):
    bad = tmp_path / "broken.xlsx"
    bad.write_text("this is not an excel file")
    result = read_workbook(bad)
    assert result.rows == []
    assert "Could not open file" in result.issues[0].message
